"""Magnet Axiom artifact catalog and Select Artifacts picker.

Catalog text comes from the Magnet Artifact Reference Guide HTM export
(stored as magnet_artifacts.json). Pictures/Videos examiner subtags stay
in this module; they are not Magnet artifact pages.
"""
import json
import tkinter as tk
from tkinter import ttk, messagebox

from ui_theme import COLORS, apply_theme, resource_dir

CATALOG_NAME = "magnet_artifacts.json"
ALL_CATEGORY = "All categories"

DEVICE_CLASS_PLATFORMS = {
    "mobile": ("Android", "iOS", "Cloud", "Refined Results"),
    "computer": ("Computer", "macOS", "Chromebook", "Cloud", "Refined Results"),
    "cloud": ("Cloud", "Refined Results"),
    "warrant": ("Cloud", "Android", "iOS", "Computer", "macOS", "Chromebook", "Refined Results"),
}

DEFAULT_PLATFORM = {
    "mobile": "Android",
    "computer": "Computer",
    "cloud": "Cloud",
    "warrant": "Cloud",
}

CATEGORY_ORDER = (
    "Operating System",
    "Application Usage",
    "Communication",
    "Email and Calendar",
    "Social Networking",
    "Web Related",
    "Location and Travel",
    "Media",
    "Documents",
    "Cloud Storage",
    "Connected Devices",
    "Peer-to-Peer",
    "Peer to Peer",
    "Encryption and Credentials",
    "Memory",
    "Volatile Artifacts",
    "Live System",
    "Additional Sources",
    "Advanced Search Tools",
    "Custom",
    "YARA Rules",
    "Refined Results",
    "Other",
)

MEDIA_PARENTS = ("Pictures", "Videos")
MEDIA_SUBTAGS = ("Child Pornography", "Child Erotica", "Age Difficult")

CUSTOM_ARTIFACTS = {
    "Pictures -- Child Pornography": {
        "description": (
            "The image files contained in this tag appear to depict child pornography. "
            "Child pornography consists of any visual depiction, including photographs, "
            "film, videos, or pictures depicting sexually explicit conduct. "
            '"Sexually explicit conduct" means material depicting any person under the '
            "age of 18 years engaged in graphic sexual intercourse, including "
            "genital-genital, oral-genital, anal-genital, or oral-anal whether between "
            "persons of the same or opposite sex, or lascivious simulated sexual "
            "intercourse where the genitals, breast or pubic area of any person is "
            "exhibited. In addition, any material depicting a minor involved in "
            "bestiality; masturbation; sadistic or masochistic abuse; or lascivious "
            "exhibition of the genitals or pubic area."
        ),
        "tag": "Pictures -- Child Pornography",
    },
    "Pictures -- Child Erotica": {
        "description": (
            "The image files contained in this tag do not meet the statutory requirement "
            "to be considered child pornography. These image files depict juvenile "
            "subjects who are shown wearing sexually suggestive clothing, posing in "
            "sexually suggestive positions, or that appear to be possessed for a sexual "
            "purpose. Image files of this nature are often referred to as "
            '"child erotica" by forensic examiners and investigators.'
        ),
        "tag": "Pictures -- Child Erotica",
    },
    "Pictures -- Age Difficult": {
        "description": (
            "The image files in this tag are pornographic in nature and depict younger "
            "looking subjects who may be juveniles under the age of 18, or may be young "
            "adults who are 18 years of age or older. Therefore, without positive "
            "identification of the subjects shown in the image files in this tag, I "
            "cannot make an accurate determination regarding the legal or illegal nature "
            "of the image files. Pornographic image files and video files of this nature "
            "are often referred to as \"age difficult\" pornography by forensic "
            "examiners and investigators."
        ),
        "tag": "Pictures -- Age Difficult",
    },
    "Videos -- Child Pornography": {
        "description": (
            "The video files contained in this tag appear to depict child pornography. "
            "Child pornography consists of any visual depiction, including photographs, "
            "film, videos, or pictures depicting sexually explicit conduct. "
            '"Sexually explicit conduct" means material depicting any person under the '
            "age of 18 years engaged in graphic sexual intercourse, including "
            "genital-genital, oral-genital, anal-genital, or oral-anal whether between "
            "persons of the same or opposite sex, or lascivious simulated sexual "
            "intercourse where the genitals, breast or pubic area of any person is "
            "exhibited. In addition, any material depicting a minor involved in "
            "bestiality; masturbation; sadistic or masochistic abuse; or lascivious "
            "exhibition of the genitals or pubic area."
        ),
        "tag": "Videos -- Child Pornography",
    },
    "Videos -- Child Erotica": {
        "description": (
            "The video files contained in this tag do not meet the statutory requirement "
            "to be considered child pornography. These video files depict juvenile "
            "subjects who are shown wearing sexually suggestive clothing, posing in "
            "sexually suggestive positions, or that appear to be possessed for a sexual "
            "purpose. Video files of this nature are often referred to as "
            '"child erotica" by forensic examiners and investigators.'
        ),
        "tag": "Videos -- Child Erotica",
    },
    "Videos -- Age Difficult": {
        "description": (
            "The video files in this tag are pornographic in nature and depict younger "
            "looking subjects who may be juveniles under the age of 18, or may be young "
            "adults who are 18 years of age or older. Therefore, without positive "
            "identification of the subjects shown in the video files in this tag, I "
            "cannot make an accurate determination regarding the legal or illegal nature "
            "of the video files. Video files of this nature are often referred to as "
            '"age difficult" pornography by forensic examiners and investigators.'
        ),
        "tag": "Videos -- Age Difficult",
    },
}

_CATALOG_CACHE = None
_INDEX_CACHE = None


def catalog_path():
    return resource_dir() / CATALOG_NAME


def _build_index(platforms):
    index = {}
    for platform, items in (platforms or {}).items():
        by_category = {}
        all_items = []
        for item in items or []:
            name = item.get("name") or ""
            if not name:
                continue
            category = item.get("category") or "Other"
            record = {
                "name": name,
                "category": category,
                "description": item.get("description") or "",
            }
            all_items.append(record)
            by_category.setdefault(category, []).append(record)
        all_items.sort(key=lambda rec: rec["name"].lower())
        for group in by_category.values():
            group.sort(key=lambda rec: rec["name"].lower())
        index[platform] = {
            "all": all_items,
            "by_category": by_category,
            "categories": sorted(by_category, key=_category_sort_key),
        }
    return index


def load_catalog():
    global _CATALOG_CACHE, _INDEX_CACHE
    if _CATALOG_CACHE is not None:
        return _CATALOG_CACHE
    path = catalog_path()
    with open(path, "r", encoding="utf-8") as handle:
        payload = json.load(handle)
    platforms = payload.get("platforms") if isinstance(payload, dict) else payload
    if not isinstance(platforms, dict):
        raise ValueError("magnet_artifacts.json is missing a platforms object")
    _CATALOG_CACHE = platforms
    _INDEX_CACHE = _build_index(platforms)
    return _CATALOG_CACHE


def load_index():
    load_catalog()
    return _INDEX_CACHE or {}


def platforms_for_device_class(device_class):
    key = (device_class or "mobile").strip().lower()
    if key in ("pc", "computer/storage", "computer", "storage"):
        key = "computer"
    if key in ("warrant", "sw", "search warrant", "warrant return"):
        key = "warrant"
    if key not in DEVICE_CLASS_PLATFORMS:
        key = "mobile"
    catalog = load_catalog()
    return [name for name in DEVICE_CLASS_PLATFORMS[key] if name in catalog]


def preferred_platform_for(device_class, device_type=None):
    platforms = platforms_for_device_class(device_class)
    hint = str(device_type or "").strip()
    if hint in platforms:
        return hint
    lowered = hint.lower()
    for name in platforms:
        if name.lower() == lowered:
            return name
    if lowered in ("ios", "iphone", "ipad", "ipod") and "iOS" in platforms:
        return "iOS"
    if lowered in ("android",) and "Android" in platforms:
        return "Android"
    if lowered in ("macos", "mac", "os x", "osx") and "macOS" in platforms:
        return "macOS"
    if lowered in ("chromebook", "chrome os") and "Chromebook" in platforms:
        return "Chromebook"
    if lowered in ("cloud",) and "Cloud" in platforms:
        return "Cloud"
    key = "computer" if (device_class or "").lower() in ("pc", "computer", "computer/storage", "storage") else (device_class or "mobile")
    fallback = DEFAULT_PLATFORM.get(key, platforms[0] if platforms else "Android")
    return fallback if fallback in platforms else (platforms[0] if platforms else "Android")


def _category_sort_key(category):
    try:
        return (0, CATEGORY_ORDER.index(category), category.lower())
    except ValueError:
        return (1, 99, category.lower())


def categories_for_platform(platform):
    info = load_index().get(platform) or {}
    return list(info.get("categories") or [])


def artifacts_for(platform, category=None, query=""):
    info = load_index().get(platform) or {}
    if category and category != ALL_CATEGORY:
        items = list(info.get("by_category", {}).get(category) or [])
    else:
        items = list(info.get("all") or [])
    needle = " ".join(str(query or "").lower().split())
    if needle:
        items = [
            item
            for item in items
            if needle in item["name"].lower() or needle in item["description"].lower()
        ]
    return items


def lookup_artifact(name, preferred_platforms=None, sources=None):
    if name in CUSTOM_ARTIFACTS:
        return dict(CUSTOM_ARTIFACTS[name])
    catalog = load_catalog()
    order = []
    if sources and name in sources:
        order.append(sources[name])
    for platform in preferred_platforms or ():
        if platform not in order:
            order.append(platform)
    for platform in catalog:
        if platform not in order:
            order.append(platform)
    for platform in order:
        for item in catalog.get(platform) or ():
            if item.get("name") == name:
                description = (item.get("description") or "").strip()
                if not description:
                    description = f"{name} is a Magnet Axiom artifact recovered from the examined evidence."
                return {"description": description, "tag": name, "platform": platform, "category": item.get("category") or ""}
    return {
        "description": f"{name} is a Magnet Axiom artifact recovered from the examined evidence.",
        "tag": name,
        "platform": "",
        "category": "",
    }


def organize_selected(selected_artifacts):
    selected = list(selected_artifacts or [])
    organized = []
    for name in selected:
        if " -- " in name or name in MEDIA_PARENTS:
            continue
        organized.append(name)
    for parent in MEDIA_PARENTS:
        if parent in selected:
            organized.append(parent)
            for name in selected:
                if name.startswith(f"{parent} -- "):
                    organized.append(name)
    for name in selected:
        if name not in organized:
            organized.append(name)
    return organized


ARTIFACT_HEADING_STYLE = "DFRArtifactHeading"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _qn(name):
    return f"{{{W_NS}}}{name}"


def _artifact_run(paragraph, text, bold=False, underline=False):
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.bold = bold
    run.font.bold = bold
    run.underline = underline
    run.font.underline = underline
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    return run


def _artifact_paragraph(doc, text, bold=False, underline=False, indent=0, heading=False):
    from docx.oxml import OxmlElement
    from docx.shared import Inches, Pt

    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Normal"]
    paragraph.paragraph_format.left_indent = Inches(indent or 0)
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    if heading:
        mark_id = str(8000 + len(list(doc.element.body.iter(_qn("bookmarkStart")))))
        ppr = paragraph._p.get_or_add_pPr()
        style = OxmlElement("w:pStyle")
        style.set(_qn("val"), ARTIFACT_HEADING_STYLE)
        existing = ppr.find(_qn("pStyle"))
        if existing is not None:
            ppr.remove(existing)
        ppr.insert(0, style)
        start = OxmlElement("w:bookmarkStart")
        start.set(_qn("id"), mark_id)
        start.set(_qn("name"), f"DFRArtifact{mark_id}")
        end = OxmlElement("w:bookmarkEnd")
        end.set(_qn("id"), mark_id)
        paragraph._p.insert(0, start)
        paragraph._p.append(end)
    _artifact_run(paragraph, text, bold=bold, underline=underline)
    return paragraph


def _paragraph_plain_text(element):
    parts = []
    for node in element.iter(_qn("t")):
        parts.append(node.text or "")
    return "".join(parts).strip()


def _run_has_tag(rpr, tag):
    if rpr is None:
        return False
    child = rpr.find(_qn(tag))
    if child is None:
        return False
    val = (child.get(_qn("val")) or "true").strip().lower()
    return val not in {"0", "false", "off"}


def is_artifact_heading_element(element):
    """True only for titles marked when the artifact list was written."""
    if element is None or not str(getattr(element, "tag", "")).endswith("}p"):
        return False
    for child in element.iter(_qn("bookmarkStart")):
        name = child.get(_qn("name")) or ""
        if name.startswith("DFRArtifact"):
            return True
    ppr = element.find(_qn("pPr"))
    if ppr is None:
        return False
    style = ppr.find(_qn("pStyle"))
    return style is not None and (style.get(_qn("val")) or "") == ARTIFACT_HEADING_STYLE


def _numbering_root(doc):
    try:
        return doc.part.numbering_part._element
    except Exception:
        return None


def _max_attr(root, tag, attr):
    highest = 0
    for child in root.findall(_qn(tag)):
        try:
            highest = max(highest, int(child.get(_qn(attr)) or 0))
        except (TypeError, ValueError):
            pass
    return highest


def ensure_artifact_numbering(doc):
    """Add a decimal numbered list to the destination document and return its numId."""
    from docx.oxml import OxmlElement

    root = _numbering_root(doc)
    if root is None:
        return None
    for abstract in root.findall(_qn("abstractNum")):
        name = abstract.find(_qn("name"))
        if name is not None and name.get(_qn("val")) == "DFR Artifact List":
            abstract_id = abstract.get(_qn("abstractNumId"))
            for num in root.findall(_qn("num")):
                pointer = num.find(_qn("abstractNumId"))
                if pointer is not None and pointer.get(_qn("val")) == abstract_id:
                    try:
                        return int(num.get(_qn("numId")))
                    except (TypeError, ValueError):
                        continue
    abstract_id = _max_attr(root, "abstractNum", "abstractNumId") + 1
    num_id = _max_attr(root, "num", "numId") + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(_qn("abstractNumId"), str(abstract_id))
    name = OxmlElement("w:name")
    name.set(_qn("val"), "DFR Artifact List")
    abstract.append(name)
    multi = OxmlElement("w:multiLevelType")
    multi.set(_qn("val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(_qn("ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(_qn("val"), "1")
    lvl.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(_qn("val"), "decimal")
    lvl.append(fmt)
    text = OxmlElement("w:lvlText")
    text.set(_qn("val"), "%1.")
    lvl.append(text)
    align = OxmlElement("w:lvlJc")
    align.set(_qn("val"), "left")
    lvl.append(align)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(_qn("left"), "720")
    ind.set(_qn("hanging"), "360")
    ppr.append(ind)
    lvl.append(ppr)
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(_qn("ascii"), "Arial")
    rfonts.set(_qn("hAnsi"), "Arial")
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(_qn("val"), "22")
    rpr.append(sz)
    lvl.append(rpr)
    abstract.append(lvl)

    first_num = root.find(_qn("num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        root.append(abstract)

    num = OxmlElement("w:num")
    num.set(_qn("numId"), str(num_id))
    pointer = OxmlElement("w:abstractNumId")
    pointer.set(_qn("val"), str(abstract_id))
    num.append(pointer)
    root.append(num)
    return num_id


def apply_artifact_list_numbering(doc):
    """Turn artifact titles into one Word numbered list on the destination document."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    if doc is None:
        return 0
    body = getattr(getattr(doc, "element", None), "body", None)
    if body is None:
        return 0
    headings = [element for element in body.iter(_qn("p")) if is_artifact_heading_element(element)]
    if not headings:
        return 0
    num_id = ensure_artifact_numbering(doc)
    if not num_id:
        return 0
    changed = 0
    for element in headings:
        ppr = element.find(_qn("pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            element.insert(0, ppr)
        for child in list(ppr):
            if child.tag == _qn("numPr"):
                ppr.remove(child)
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        nid = OxmlElement("w:numId")
        nid.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(nid)
        ppr.append(num_pr)
        ind = ppr.find(_qn("ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            ppr.append(ind)
        ind.set(_qn("left"), "720")
        ind.set(_qn("hanging"), "360")
        spacing = ppr.find(_qn("spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            ppr.append(spacing)
        spacing.set(_qn("before"), "0")
        spacing.set(_qn("after"), "0")
        spacing.set(_qn("line"), "240")
        spacing.set(_qn("lineRule"), "auto")
        changed += 1
    return changed


def _tags_for_artifact(artifact, info, report_tags):
    labels = []
    seen = set()

    def add(label):
        text = (label or "").strip()
        if not text:
            return
        key = " ".join(text.replace("_", "/").split()).lower()
        if key in seen:
            return
        seen.add(key)
        labels.append(text)

    if report_tags:
        for key in (artifact, info.get("tag")):
            for label in report_tags.get(key) or ():
                add(label)
    if not labels:
        # A loaded AXIOM report should not invent a generic Pictures/Videos
        # tag for items that only have CP / CE / AD tags.
        if report_tags and artifact in ("Pictures", "Videos"):
            return labels
        add(info.get("tag") or artifact)
    return labels


def _tag_item_count(label, tag_counts):
    if not tag_counts or not label:
        return None
    if label in tag_counts:
        try:
            return int(tag_counts[label])
        except (TypeError, ValueError):
            return None
    needle = " ".join(label.replace("_", "/").split()).lower()
    for key, value in tag_counts.items():
        if " ".join(str(key).replace("_", "/").split()).lower() == needle:
            try:
                return int(value)
            except (TypeError, ValueError):
                return None
    return None


def _blank_line(doc):
    doc.add_paragraph()


def _write_device_artifacts_block(doc, label, tag_counts, indent=0):
    count = _tag_item_count(label, tag_counts)
    if count is not None:
        noun = "artifact" if count == 1 else "artifacts"
        _artifact_paragraph(doc, f"This tag contains {count} {noun}.", indent=0)
        _blank_line(doc)
    _artifact_paragraph(doc, "(DEVICE ARTIFACTS)", indent=0)
    _blank_line(doc)
    _blank_line(doc)


def write_artifact_paragraphs(doc, selected_artifacts, style="mobile", preferred_platforms=None, sources=None, report_tags=None, tag_counts=None):
    organized = organize_selected(selected_artifacts)
    if not organized:
        return
    for artifact in organized:
        info = lookup_artifact(artifact, preferred_platforms=preferred_platforms, sources=sources)
        is_media_subtag = artifact.startswith("Pictures -- ") or artifact.startswith("Videos -- ")
        tag_labels = _tags_for_artifact(artifact, info, report_tags)
        if not is_media_subtag:
            _artifact_paragraph(doc, artifact.upper(), bold=True, heading=True)
            _blank_line(doc)
            _artifact_paragraph(doc, info["description"], indent=0)
            _blank_line(doc)
            for label in tag_labels:
                _artifact_paragraph(doc, f"Tag: {label}", underline=True, indent=0)
                _blank_line(doc)
                _write_device_artifacts_block(doc, label, tag_counts)
        else:
            first = True
            for label in tag_labels:
                _artifact_paragraph(doc, f"Tag: {label}", underline=True, indent=0)
                _blank_line(doc)
                if first:
                    _artifact_paragraph(doc, info["description"], indent=0)
                    _blank_line(doc)
                    first = False
                _write_device_artifacts_block(doc, label, tag_counts)


def open_artifact_picker(parent, device_class="mobile", device_type=None, previously_selected=None, sources=None):
    """Modal picker. Returns (selected_names, sources_by_name)."""
    try:
        load_catalog()
    except Exception as exc:
        messagebox.showerror(
            "Artifact Catalog",
            f"Could not load Magnet artifact descriptions:\n{exc}",
        )
        return list(previously_selected or []), dict(sources or {})

    platforms = platforms_for_device_class(device_class)
    if not platforms:
        messagebox.showerror("Artifact Catalog", "No Magnet artifacts are available for this device type.")
        return list(previously_selected or []), dict(sources or {})

    selected = set(previously_selected or [])
    source_map = dict(sources or {})
    start_platform = preferred_platform_for(device_class, device_type)

    window = tk.Toplevel(parent)
    raw_class = (device_class or "mobile").lower()
    if raw_class in ("pc", "computer/storage", "storage"):
        raw_class = "computer"
    if raw_class in ("sw", "search warrant", "warrant return"):
        raw_class = "warrant"
    title_class = {
        "mobile": "Mobile",
        "computer": "Computer / Storage",
        "cloud": "Cloud",
        "warrant": "Warrant / Cloud",
    }.get(raw_class, "Mobile")
    window.title(f"Select {title_class} Artifacts")
    apply_theme(window)
    window.transient(parent)
    window.grab_set()

    screen_w = window.winfo_screenwidth()
    screen_h = window.winfo_screenheight()
    width = min(980, max(820, screen_w - 80))
    height = min(680, max(560, screen_h - 80))
    x = max(0, (screen_w - width) // 2)
    y = max(20, (screen_h - height) // 2)
    window.geometry(f"{width}x{height}+{x}+{y}")
    window.minsize(780, 520)

    result = {"names": list(previously_selected or []), "sources": dict(source_map)}

    main = ttk.Frame(window, padding=10)
    main.pack(fill="both", expand=True)
    main.columnconfigure(1, weight=1)
    main.rowconfigure(1, weight=1)

    header = ttk.Frame(main)
    header.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 8))
    ttk.Label(
        header,
        text="Choose a source on the left, then check the Magnet artifacts tagged in this examination.",
        wraplength=860,
    ).pack(anchor="w")

    side = ttk.Frame(main, style="Card.TFrame", padding=8)
    side.grid(row=1, column=0, sticky="nsw", padx=(0, 10))
    side.rowconfigure(3, weight=1)

    ttk.Label(side, text="Source", style="CategoryHeader.TLabel").grid(row=0, column=0, sticky="w", pady=(0, 4))
    platform_var = tk.StringVar(window, value=start_platform if start_platform in platforms else platforms[0])
    platform_box = ttk.Frame(side, style="Card.TFrame")
    platform_box.grid(row=1, column=0, sticky="ew", pady=(0, 10))
    for name in platforms:
        ttk.Radiobutton(
            platform_box,
            text=name,
            value=name,
            variable=platform_var,
        ).pack(anchor="w", pady=1)

    ttk.Label(side, text="Category", style="CategoryHeader.TLabel").grid(row=2, column=0, sticky="w", pady=(0, 4))
    category_list = tk.Listbox(
        side,
        height=18,
        width=28,
        exportselection=False,
        bg=COLORS["entry_bg"],
        fg=COLORS["text"],
        selectbackground=COLORS["accent_dark"],
        selectforeground=COLORS["button_fg"],
        highlightthickness=1,
        highlightbackground=COLORS["border"],
        relief="flat",
        activestyle="none",
    )
    category_list.grid(row=3, column=0, sticky="nsw")
    cat_scroll = ttk.Scrollbar(side, orient="vertical", command=category_list.yview)
    cat_scroll.grid(row=3, column=1, sticky="ns")
    category_list.configure(yscrollcommand=cat_scroll.set)

    right = ttk.Frame(main)
    right.grid(row=1, column=1, sticky="nsew")
    right.columnconfigure(0, weight=1)
    right.rowconfigure(1, weight=1)

    tools = ttk.Frame(right)
    tools.grid(row=0, column=0, sticky="ew", pady=(0, 6))
    tools.columnconfigure(1, weight=1)
    ttk.Label(tools, text="Search:").grid(row=0, column=0, sticky="w")
    search_var = tk.StringVar(window)
    search_entry = ttk.Entry(tools, textvariable=search_var)
    search_entry.grid(row=0, column=1, sticky="ew", padx=6)
    count_var = tk.StringVar(window, value="")
    ttk.Label(tools, textvariable=count_var, style="Hint.TLabel").grid(row=0, column=2, sticky="e")

    list_holder = ttk.Frame(right, style="Card.TFrame")
    list_holder.grid(row=1, column=0, sticky="nsew")
    list_holder.columnconfigure(0, weight=1)
    list_holder.rowconfigure(0, weight=1)
    item_list = tk.Listbox(
        list_holder,
        activestyle="none",
        exportselection=False,
        bg=COLORS["panel"],
        fg=COLORS["text"],
        selectbackground=COLORS["card"],
        selectforeground=COLORS["text"],
        highlightthickness=1,
        highlightbackground=COLORS["border"],
        highlightcolor=COLORS["border"],
        relief="flat",
        font=("Segoe UI", 10),
    )
    item_scroll = ttk.Scrollbar(list_holder, orient="vertical", command=item_list.yview)
    item_list.grid(row=0, column=0, sticky="nsew")
    item_scroll.grid(row=0, column=1, sticky="ns")
    item_list.configure(yscrollcommand=item_scroll.set)

    visible_rows = []
    search_job = {"id": None}

    def _mark(checked):
        return "☑  " if checked else "☐  "

    def _row_text(name, kind="main"):
        checked = name in selected
        if kind == "sub":
            label = name.split(" -- ", 1)[-1]
            prefix = "     ☑  • " if checked else "     ☐  • "
            return prefix + label
        return _mark(checked) + name

    def _selected_count():
        return len(selected)

    def refresh_count():
        platform = platform_var.get()
        category = _current_category()
        count_var.set(f"{_selected_count()} selected  ·  {len(visible_rows)} listed")
        window.title(f"Select {title_class} Artifacts — {platform} / {category}")

    def _current_category():
        choices = [ALL_CATEGORY] + categories_for_platform(platform_var.get())
        try:
            index = category_list.curselection()[0]
            return choices[index]
        except Exception:
            return choices[1] if len(choices) > 1 else ALL_CATEGORY

    def rebuild_categories(keep_name=None):
        cats = [ALL_CATEGORY] + categories_for_platform(platform_var.get())
        category_list.delete(0, tk.END)
        for cat in cats:
            category_list.insert(tk.END, cat)
        if keep_name in cats:
            target = keep_name
        elif len(cats) > 1:
            target = cats[1]
        else:
            target = ALL_CATEGORY
        category_list.selection_clear(0, tk.END)
        category_list.selection_set(cats.index(target))
        category_list.see(cats.index(target))

    def toggle_name(name, value, platform):
        if value:
            selected.add(name)
            source_map[name] = platform
        else:
            selected.discard(name)
            source_map.pop(name, None)
            if name in MEDIA_PARENTS:
                for sub in list(selected):
                    if sub.startswith(f"{name} -- "):
                        selected.discard(sub)
                        source_map.pop(sub, None)

    def render_items(keep_index=None):
        platform = platform_var.get()
        category = _current_category()
        items = artifacts_for(platform, category, search_var.get())
        rows = []
        for item in items:
            name = item["name"]
            rows.append({"name": name, "kind": "main"})
            if name in MEDIA_PARENTS and name in selected:
                for label in MEDIA_SUBTAGS:
                    rows.append({"name": f"{name} -- {label}", "kind": "sub"})
        visible_rows[:] = rows
        item_list.delete(0, tk.END)
        if not rows:
            item_list.insert(tk.END, "  No artifacts match this source, category, and search.")
            refresh_count()
            return
        for row in rows:
            item_list.insert(tk.END, _row_text(row["name"], row["kind"]))
        if keep_index is not None and rows:
            item_list.selection_clear(0, tk.END)
            item_list.selection_set(min(keep_index, len(rows) - 1))
            item_list.activate(min(keep_index, len(rows) - 1))
        refresh_count()

    def _toggle_index(index):
        if index < 0 or index >= len(visible_rows):
            return
        row = visible_rows[index]
        name = row["name"]
        platform = platform_var.get()
        was_checked = name in selected
        toggle_name(name, not was_checked, platform)
        if name in MEDIA_PARENTS:
            render_items(keep_index=index)
        else:
            item_list.delete(index)
            item_list.insert(index, _row_text(name, row["kind"]))
            item_list.selection_clear(0, tk.END)
            item_list.selection_set(index)
            item_list.activate(index)
            refresh_count()

    def on_item_click(event):
        index = item_list.nearest(event.y)
        _toggle_index(index)

    def on_item_key(event):
        if event.keysym in ("space", "Return"):
            try:
                _toggle_index(item_list.curselection()[0])
            except Exception:
                pass
            return "break"

    def on_platform_change(*_args):
        current = _current_category()
        rebuild_categories(keep_name=current)
        render_items()

    def on_category_change(_event=None):
        render_items()

    def on_search(*_args):
        job = search_job.get("id")
        if job:
            window.after_cancel(job)
        search_job["id"] = window.after(120, render_items)

    def select_visible():
        platform = platform_var.get()
        for row in visible_rows:
            name = row["name"]
            if row["kind"] == "sub":
                parent = name.split(" -- ", 1)[0]
                if parent not in selected:
                    continue
            selected.add(name)
            source_map[name] = platform
        render_items()

    def deselect_visible():
        for row in list(visible_rows):
            name = row["name"]
            selected.discard(name)
            source_map.pop(name, None)
            if name in MEDIA_PARENTS:
                for sub in list(selected):
                    if sub.startswith(f"{name} -- "):
                        selected.discard(sub)
                        source_map.pop(sub, None)
        render_items()

    def clear_all():
        selected.clear()
        source_map.clear()
        render_items()

    def finish():
        ordered = organize_selected(list(selected))
        result["names"] = ordered
        result["sources"] = {name: source_map[name] for name in ordered if name in source_map}
        window.destroy()

    def cancel():
        window.destroy()

    buttons = ttk.Frame(main)
    buttons.grid(row=2, column=0, columnspan=2, sticky="ew", pady=(10, 0))
    ttk.Button(buttons, text="Select listed", command=select_visible).pack(side="left", padx=4)
    ttk.Button(buttons, text="Clear listed", command=deselect_visible).pack(side="left", padx=4)
    ttk.Button(buttons, text="Clear all", command=clear_all).pack(side="left", padx=4)
    ttk.Button(buttons, text="Done", command=finish).pack(side="right", padx=4)
    ttk.Button(buttons, text="Cancel", command=cancel).pack(side="right", padx=4)

    platform_var.trace_add("write", on_platform_change)
    search_var.trace_add("write", on_search)
    category_list.bind("<<ListboxSelect>>", on_category_change)
    item_list.bind("<Button-1>", on_item_click)
    item_list.bind("<space>", on_item_key)
    item_list.bind("<Return>", on_item_key)
    rebuild_categories()
    render_items()
    search_entry.focus_set()
    window.protocol("WM_DELETE_WINDOW", cancel)
    parent.wait_window(window)
    job = search_job.get("id")
    if job:
        try:
            window.after_cancel(job)
        except Exception:
            pass
    return result["names"], result["sources"]


def pick_artifacts(parent, device_class="mobile", device_type=None):
    previous = getattr(parent, "selected_artifacts", []) or []
    sources = getattr(parent, "selected_artifact_sources", {}) or {}
    names, source_map = open_artifact_picker(
        parent,
        device_class=device_class,
        device_type=device_type,
        previously_selected=previous,
        sources=sources,
    )
    parent.selected_artifacts = names
    parent.selected_artifact_sources = source_map
    return names
