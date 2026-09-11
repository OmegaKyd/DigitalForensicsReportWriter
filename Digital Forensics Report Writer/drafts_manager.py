"""Unfinished-report drafts: save, list, load, and delete after generate."""
from __future__ import annotations

import json
import os
import re
import sys
from datetime import datetime
from pathlib import Path

from ui_theme import APP_NAME, APP_VERSION, COLORS

REPORT_TYPES = {
    "mobile_full": "Mobile — Full Exam",
    "mobile_portable": "Mobile — Portable Case",
    "pc_full": "PC — Full Exam",
    "pc_portable": "PC — Portable Case",
    "warrant": "Warrant Data Returns",
}

WIDGET_ATTRS = (
    "role_type",
    "request_date",
    "request_agency",
    "request_title_type",
    "request_title_entry",
    "request_officer",
    "case_type",
    "offense_type",
    "legal_self",
    "legal_authority",
    "sw_service_date",
    "transfer_title",
    "transfer_officer",
    "transfer_agency",
    "transfer_date",
    "examiner_agency_type",
    "examiner_agency_entry",
    "examiner_title_type",
    "examiner_title_entry",
    "examiner_name",
    "dfr_number",
    "case_number",
    "evidence_number",
    "device_owner",
    "account_owner",
    "device_color",
    "device_capacity",
    "device_iccid",
    "device_passcode",
    "device_password",
    "device_carrier",
    "device_model",
    "device_type",
    "output_filename",
    "save_location",
    "service_provider",
    "account_identifier",
    "account_data_size",
    "warrant_service_date",
    "data_return_date",
    "time_frame_start_date",
    "time_frame_end_date",
    "time_frame_start",
    "time_frame_end",
    "case_time_frame_start",
    "case_time_frame_end",
    "case_time_frame_start_date",
    "case_time_frame_end_date",
    "case_agent_time_frame_start",
    "case_agent_time_frame_end",
    "primary_case_offense",
    "airplane_mode",
)

VAR_ATTRS = (
    "axiom_var",
    "cellebrite_var",
    "griffeye_var",
    "cb_axiom_var",
    "cb_cellebrite_var",
    "cb_griffeye_var",
    "cb_manual_var",
    "device_transfer_var",
    "no_evidence_var",
    "time_frame_limited_var",
    "time_frame_var",
    "case_time_frame_var",
    "case_time_frame_limited_var",
    "case_agent_time_frame_var",
    "xways_var",
)

STATE_ATTRS = (
    "selected_artifacts",
    "selected_artifact_sources",
    "axiom_report_tags",
    "axiom_tag_counts",
    "axiom_version",
    "axiom_report_path",
    "extraction_file",
    "template_file",
    "multiple_extractions",
    "extraction_type",
)

TOGGLE_METHODS = (
    "toggle_request_content",
    "toggle_transfer_fields",
    "toggle_title_entry",
    "toggle_agency_entry",
    "toggle_examiner_title_entry",
    "toggle_examiner_title_other",
    "toggle_requesting_officer_title_other",
    "toggle_sw_date",
    "toggle_sw_date_and_time_frame",
    "toggle_time_frame_section",
    "toggle_time_frame_fields",
    "toggle_case_time_frame_section",
    "toggle_case_time_frame_fields",
    "toggle_case_agent_time_frame_fields",
    "toggle_device_type_fields",
    "on_forensic_software_change",
    "on_processing_software_toggled",
    "_sync_axiom_artifact_controls",
    "toggle_artifacts_button",
    "update_artifacts_count_label",
)


def drafts_dir(create=True):
    if getattr(sys, "frozen", False):
        if sys.platform == "win32":
            base = os.environ.get("APPDATA") or str(Path.home() / "AppData" / "Roaming")
        else:
            base = os.environ.get("XDG_CONFIG_HOME") or str(Path.home() / ".config")
        path = Path(base) / APP_NAME / "Unfinished Reports"
    else:
        path = Path(__file__).resolve().parent / "Unfinished Reports"
    if create:
        path.mkdir(parents=True, exist_ok=True)
    return path


def _safe_token(value):
    text = re.sub(r"[^A-Za-z0-9._-]+", "_", str(value or "").strip())
    return text.strip("._-") or "untitled"


def draft_filename(report_type, dfr_number):
    return f"{_safe_token(report_type)}__{_safe_token(dfr_number)}.json"


def form_dfr_number(app):
    widget = getattr(app, "dfr_number", None)
    if widget is None:
        return ""
    try:
        return widget.get().strip()
    except Exception:
        return ""


def form_owner_name(app):
    for name in ("device_owner", "account_owner"):
        widget = getattr(app, name, None)
        if widget is None:
            continue
        try:
            value = widget.get().strip()
        except Exception:
            value = ""
        if value:
            return value
    return ""


def draft_label(dfr_number, owner):
    dfr = (dfr_number or "").strip() or "Unnumbered"
    owner = (owner or "").strip()
    if owner:
        return f"{dfr} - {owner}"
    return dfr


def _widget_value(widget):
    try:
        return widget.get()
    except Exception:
        return None


def _set_widget_value(widget, value):
    if value is None:
        return
    text = "" if value is None else str(value)
    try:
        widget.delete(0, "end")
        widget.insert(0, text)
        return
    except Exception:
        pass
    try:
        widget.set(text)
    except Exception:
        pass


def collect_draft(app):
    report_type = getattr(app, "report_type", "") or ""
    dfr = form_dfr_number(app)
    owner = form_owner_name(app)
    widgets = {}
    for name in WIDGET_ATTRS:
        widget = getattr(app, name, None)
        if widget is None:
            continue
        value = _widget_value(widget)
        if value is None:
            continue
        widgets[name] = value
    variables = {}
    for name in VAR_ATTRS:
        var = getattr(app, name, None)
        if var is None:
            continue
        try:
            variables[name] = var.get()
        except Exception:
            pass
    notes = exam_notes_text(app)
    state = {}
    for name in STATE_ATTRS:
        if hasattr(app, name):
            value = getattr(app, name)
            if callable(value):
                continue
            state[name] = value
    choice = getattr(app, "template_choice", None)
    if choice is not None:
        try:
            state["template_choice"] = choice.get()
        except Exception:
            pass
    preview = getattr(app, "_preview_overrides", None)
    if isinstance(preview, dict):
        state["preview_overrides"] = preview
    extracted = getattr(app, "_preview_extracted", None)
    if isinstance(extracted, dict):
        state["preview_extracted"] = extracted
    checklist = collect_checklist_state(app)
    if checklist:
        state["checklist"] = checklist
    return {
        "version": 1,
        "app_version": APP_VERSION,
        "report_type": report_type,
        "dfr_number": dfr,
        "owner": owner,
        "label": draft_label(dfr, owner),
        "updated": datetime.now().isoformat(timespec="seconds"),
        "widgets": widgets,
        "vars": variables,
        "notes": notes,
        "state": state,
    }


def save_progress(app):
    from report_common import is_complete_dfr_number

    report_type = getattr(app, "report_type", "") or ""
    if report_type not in REPORT_TYPES:
        raise ValueError("This window cannot save unfinished report progress.")
    dfr = form_dfr_number(app)
    if not is_complete_dfr_number(dfr):
        raise ValueError("Enter a complete Report Number before saving progress.")
    payload = collect_draft(app)
    path = drafts_dir() / draft_filename(report_type, dfr)
    previous = getattr(app, "_draft_path", None)
    with open(path, "w", encoding="utf-8") as handle:
        json.dump(payload, handle, indent=2, ensure_ascii=False)
    if previous and Path(previous).resolve() != path.resolve():
        delete_draft_file(previous)
    app._draft_path = str(path)
    return path, payload["label"]


def list_drafts():
    items = []
    folder = drafts_dir(create=False)
    try:
        files = list(folder.glob("*.json")) if folder.is_dir() else []
    except Exception:
        files = []
    for path in files:
        try:
            with open(path, "r", encoding="utf-8") as handle:
                data = json.load(handle)
        except Exception:
            continue
        if not isinstance(data, dict):
            continue
        report_type = data.get("report_type") or ""
        if report_type not in REPORT_TYPES:
            continue
        label = data.get("label") or draft_label(data.get("dfr_number"), data.get("owner"))
        updated = data.get("updated") or ""
        items.append({
            "path": path,
            "label": label,
            "report_type": report_type,
            "report_label": REPORT_TYPES.get(report_type, report_type),
            "updated": updated,
            "data": data,
        })
    items.sort(key=lambda item: (item.get("updated") or "", item["label"]), reverse=True)
    return items


def load_draft_file(path):
    with open(path, "r", encoding="utf-8") as handle:
        data = json.load(handle)
    if not isinstance(data, dict):
        raise ValueError("Invalid unfinished report file.")
    return data


def delete_draft_file(path):
    try:
        target = Path(path)
        if target.is_file():
            target.unlink()
            return True
    except Exception:
        return False
    return False


def delete_draft_for_app(app):
    report_type = getattr(app, "report_type", "") or ""
    dfr = form_dfr_number(app)
    removed = False
    stored = getattr(app, "_draft_path", None)
    if stored:
        removed = delete_draft_file(stored) or removed
    if report_type and dfr:
        removed = delete_draft_file(drafts_dir() / draft_filename(report_type, dfr)) or removed
    app._draft_path = None
    return removed


def exam_notes_text(app):
    widget = getattr(app, "exam_notes", None)
    if widget is None:
        return ""
    try:
        return widget.get("1.0", "end-1c").strip()
    except Exception:
        return ""


def set_exam_notes(app, text):
    widget = getattr(app, "exam_notes", None)
    if widget is None:
        return
    try:
        widget.delete("1.0", "end")
        if text:
            widget.insert("1.0", text)
    except Exception:
        pass


def prepend_exam_notes(doc, app):
    """Write optional examiner notes at the front of the PY_TEXT document."""
    from docx.shared import Pt

    text = exam_notes_text(app)
    if not text:
        return
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    wrote = False
    for line in lines:
        if not line.strip() and not wrote:
            continue
        paragraph = doc.add_paragraph(line)
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(11)
        wrote = True
    if wrote:
        divider = doc.add_paragraph("***************")
        for run in divider.runs:
            run.font.name = "Arial"
            run.font.size = Pt(11)
        spacer = doc.add_paragraph("")
        for run in spacer.runs:
            run.font.name = "Arial"
            run.font.size = Pt(11)


ANDROID_CHECKLIST_ITEMS = (
    "Enabled Developer Options",
    "Enabled USB Debugging",
    "Set to Stay Awake",
    "USB Settings Set to File Transfer or MTP",
    "Verify Apps Turned Off",
    "Screen Timeout Set to Longest Setting",
    "Unknown Sources Selected in Developer Options",
    "Turned Lock Off",
    "Enabled Recovery Mode",
)

IOS_CHECKLIST_ITEMS = (
    "Enabled Developer Options",
    "Trust Computer",
    "Screen Timeout Set to Never",
    "Turned Off Low Power Mode",
    "Turned Off Lock",
    "Entered Recovery Mode",
    "Entered DFU Mode",
)

CHECK_MARK = "☑"
UNCHECK_MARK = "☐"


def add_notes_tab(app, form_tabs, mobile_checklists=False):
    import tkinter as tk
    from tkinter import ttk

    page = form_tabs.add_tab("notes", "Notes")
    frame = ttk.LabelFrame(page, text="Examination Notes (optional)", padding="10")
    frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
    ttk.Label(
        frame,
        text=(
            "Optional. On mobile reports, notes print after the device checklist at PY_CHECKLIST. "
            "On other reports, notes print at the start of PY_TEXT."
            if mobile_checklists
            else "Optional. If used, these notes print at the start of the report narrative, before the regular paragraphs."
        ),
        style="Hint.TLabel",
        wraplength=640,
    ).pack(anchor="w")
    wrap = ttk.Frame(frame)
    wrap.pack(fill=tk.BOTH, expand=True, pady=(8, 0))
    text = tk.Text(
        wrap,
        wrap="word",
        height=10 if mobile_checklists else 18,
        bg=COLORS["entry_bg"],
        fg=COLORS["text"],
        insertbackground=COLORS["accent"],
        relief="flat",
        font=("Segoe UI", 10),
    )
    scroll = ttk.Scrollbar(wrap, orient=tk.VERTICAL, command=text.yview)
    text.configure(yscrollcommand=scroll.set)
    text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
    scroll.pack(side=tk.RIGHT, fill=tk.Y)
    app.exam_notes = text
    app.tab_notes = page
    if mobile_checklists:
        add_mobile_checklists(app, page)
    return page


def add_mobile_checklists(app, parent):
    import tkinter as tk
    from tkinter import ttk, messagebox

    box = ttk.LabelFrame(parent, text="Device Prep Checklists (optional)", padding="10")
    box.pack(fill=tk.X, expand=False, padx=5, pady=(0, 5))
    ttk.Label(
        box,
        text="Check items from only one list. That list prints at PY_CHECKLIST, with a checked box on selected items.",
        style="Hint.TLabel",
        wraplength=640,
    ).pack(anchor="w", pady=(0, 8))

    columns = ttk.Frame(box)
    columns.pack(fill=tk.X, expand=True)
    columns.columnconfigure(0, weight=1)
    columns.columnconfigure(1, weight=1)

    def build_column(parent_frame, column, title, items, attr_name):
        col = ttk.Frame(parent_frame)
        col.grid(row=0, column=column, sticky="nsew", padx=(0, 12) if column == 0 else (12, 0))
        tk.Label(
            col,
            text=title,
            bg=COLORS["bg"],
            fg=COLORS["text"],
            font=("Segoe UI", 10, "bold", "underline"),
        ).pack(anchor="w", pady=(0, 6))
        variables = []
        for item in items:
            var = tk.IntVar(app, value=0)
            ttk.Checkbutton(
                col,
                text=item,
                variable=var,
                onvalue=1,
                offvalue=0,
                command=lambda kind=attr_name, current=var: _enforce_single_checklist(app, kind, current),
            ).pack(anchor="w", pady=1)
            variables.append(var)
        setattr(app, attr_name, variables)
        return variables

    build_column(columns, 0, "ANDROID CHECKLIST", ANDROID_CHECKLIST_ITEMS, "android_checklist_vars")
    build_column(columns, 1, "iOS CHECKLIST", IOS_CHECKLIST_ITEMS, "ios_checklist_vars")


def _enforce_single_checklist(app, changed_attr, current_var):
    from tkinter import messagebox

    if getattr(app, "_restoring_draft", False):
        return
    if not current_var.get():
        return
    other_attr = "ios_checklist_vars" if changed_attr == "android_checklist_vars" else "android_checklist_vars"
    other_vars = getattr(app, other_attr, []) or []
    if any(var.get() for var in other_vars):
        current_var.set(0)
        messagebox.showinfo(
            "Device Prep Checklist",
            "Select items from only one checklist (Android or iOS).",
            parent=app,
        )


def collect_checklist_state(app):
    def values(attr):
        return [int(var.get()) for var in (getattr(app, attr, None) or [])]

    android = values("android_checklist_vars")
    ios = values("ios_checklist_vars")
    if not android and not ios:
        return None
    return {"android": android, "ios": ios}


def apply_checklist_state(app, state):
    if not isinstance(state, dict):
        return
    already = getattr(app, "_restoring_draft", False)
    app._restoring_draft = True
    try:
        for attr, key in (("android_checklist_vars", "android"), ("ios_checklist_vars", "ios")):
            variables = getattr(app, attr, None) or []
            saved = state.get(key) or []
            for index, var in enumerate(variables):
                try:
                    var.set(int(saved[index]) if index < len(saved) else 0)
                except Exception:
                    var.set(0)
    finally:
        if not already:
            app._restoring_draft = False


def active_checklist(app):
    android_vars = getattr(app, "android_checklist_vars", None) or []
    ios_vars = getattr(app, "ios_checklist_vars", None) or []
    android_on = any(var.get() for var in android_vars)
    ios_on = any(var.get() for var in ios_vars)
    if android_on and not ios_on:
        return "ANDROID CHECKLIST", ANDROID_CHECKLIST_ITEMS, android_vars
    if ios_on and not android_on:
        return "iOS CHECKLIST", IOS_CHECKLIST_ITEMS, ios_vars
    return None, (), ()


def _add_arial_run(paragraph, text, bold=False, underline=False):
    from docx.shared import Pt

    run = paragraph.add_run(text)
    run.bold = bold
    run.underline = underline
    run.font.name = "Arial"
    run.font.size = Pt(11)
    return run


def build_checklist_document(app):
    from docx import Document

    doc = Document()
    title, items, variables = active_checklist(app)
    notes = exam_notes_text(app)
    if title:
        header = doc.add_paragraph()
        _add_arial_run(header, title, bold=True, underline=True)
        for item, var in zip(items, variables):
            mark = CHECK_MARK if var.get() else UNCHECK_MARK
            line = doc.add_paragraph()
            _add_arial_run(line, f"{mark} {item}")
    if title and notes:
        divider = doc.add_paragraph()
        _add_arial_run(divider, "***************")
    if notes:
        lines = notes.replace("\r\n", "\n").replace("\r", "\n").split("\n")
        wrote = False
        for line in lines:
            if not line.strip() and not wrote:
                continue
            paragraph = doc.add_paragraph()
            _add_arial_run(paragraph, line)
            wrote = True
    if not title and not notes:
        doc.add_paragraph("")
    return doc


def attach_checklist_document(search_docs, app):
    search_docs["PY_CHECKLIST"] = build_checklist_document(app)
    return search_docs["PY_CHECKLIST"]


def replace_py_checklist(doc, app):
    checklist_doc = build_checklist_document(app)
    title, _items, _variables = active_checklist(app)
    notes = exam_notes_text(app)
    for para in doc.paragraphs:
        if "PY_CHECKLIST" not in para.text:
            continue
        parent = para._element.getparent()
        insert_index = parent.index(para._element)
        if not title and not notes:
            para.clear()
            return True
        for new_para in reversed(checklist_doc.paragraphs):
            copied = doc.add_paragraph()
            for run in new_para.runs:
                new_run = copied.add_run(run.text)
                new_run.font.bold = run.font.bold
                new_run.font.italic = run.font.italic
                new_run.font.underline = run.font.underline
                if run.font.name:
                    new_run.font.name = run.font.name
                if run.font.size:
                    new_run.font.size = run.font.size
            parent.insert(insert_index, copied._element)
        parent.remove(para._element)
        return True
    return False


def is_restoring_draft(app):
    return bool(getattr(app, "_restoring_draft", False))


def _apply_draft_fields(app, payload):
    widgets = payload.get("widgets") or {}
    for name, value in widgets.items():
        widget = getattr(app, name, None)
        if widget is None:
            continue
        _set_widget_value(widget, value)
    variables = payload.get("vars") or {}
    for name, value in variables.items():
        var = getattr(app, name, None)
        if var is None:
            continue
        try:
            var.set(value)
        except Exception:
            pass
    set_exam_notes(app, payload.get("notes") or "")


def apply_draft(app, payload):
    if not payload:
        return
    app._restoring_draft = True
    state = payload.get("state") or {}
    try:
        _apply_draft_fields(app, payload)
        for name in STATE_ATTRS:
            if name in state:
                setattr(app, name, state[name])
        if "preview_overrides" in state:
            app._preview_overrides = dict(state.get("preview_overrides") or {})
        if "preview_extracted" in state:
            app._preview_extracted = dict(state.get("preview_extracted") or {})
        if "checklist" in state:
            apply_checklist_state(app, state.get("checklist"))
        choice = state.get("template_choice")
        combo = getattr(app, "template_choice", None)
        if choice and combo is not None:
            try:
                combo.set(choice)
                choices = getattr(app, "_template_choices", {}) or {}
                if choice in choices:
                    app.template_file = choices[choice]
            except Exception:
                pass
        template_file = state.get("template_file")
        if template_file and Path(template_file).is_file():
            app.template_file = template_file
            try:
                from report_common import sync_template_choice
                sync_template_choice(app, template_file)
            except Exception:
                pass
        extraction = state.get("extraction_file") or ""
        label = getattr(app, "extraction_drop_label", None)
        if label is not None and extraction:
            exists = os.path.isfile(extraction)
            name = os.path.basename(extraction)
            try:
                label.configure(text=f"Selected: {name}" + ("" if exists else " (missing)"))
            except Exception:
                pass
        axiom_path = state.get("axiom_report_path") or ""
        axiom_label = getattr(app, "axiom_report_drop_label", None)
        if axiom_label is not None and axiom_path:
            try:
                axiom_label.configure(text=f"AXIOM report: {os.path.basename(axiom_path)}")
            except Exception:
                pass
        for method_name in TOGGLE_METHODS:
            method = getattr(app, method_name, None)
            if callable(method):
                try:
                    method()
                except Exception:
                    pass
        # Role-change toggles wipe Request Info; put saved values back after layout.
        _apply_draft_fields(app, payload)
        for name in ("selected_artifacts", "selected_artifact_sources", "axiom_report_tags", "axiom_tag_counts", "axiom_version", "axiom_report_path"):
            if name in state:
                setattr(app, name, state[name])
        count = getattr(app, "update_artifacts_count_label", None)
        if callable(count):
            try:
                count()
            except Exception:
                pass
        refresh = getattr(app, "refresh_tab_status", None)
        if callable(refresh):
            try:
                refresh()
            except Exception:
                pass
        app._draft_path = payload.get("_path") or getattr(app, "_draft_path", None)
    finally:
        app._restoring_draft = False


def launch_report_type(start, report_type, draft=None):
    factories = {
        "mobile_full": start.launch_mobile_full_exam,
        "mobile_portable": start.launch_mobile_portable_case,
        "pc_full": start.launch_pc_full_exam,
        "pc_portable": start.launch_pc_portable_case,
        "warrant": start.launch_warrant_data_returns,
    }
    factory = factories.get(report_type)
    if factory is None:
        raise ValueError(f"Unknown report type: {report_type}")
    start._pending_draft = draft
    factory()


def apply_prefix_to_open_forms(root, old_prefix, new_prefix):
    if root is None or not new_prefix:
        return
    seen = set()
    stack = [root]
    master = getattr(root, "master", None)
    if master is not None:
        stack.append(master)
    open_app = getattr(root, "open_app", None)
    if open_app is not None:
        stack.append(open_app)
    while stack:
        widget = stack.pop()
        ident = id(widget)
        if ident in seen:
            continue
        seen.add(ident)
        field = getattr(widget, "dfr_number", None)
        if field is not None:
            try:
                current = field.get().strip()
            except Exception:
                current = ""
            if not current or current == old_prefix or current == (old_prefix or "").rstrip("-"):
                try:
                    field.delete(0, "end")
                    field.insert(0, new_prefix)
                except Exception:
                    pass
        try:
            stack.extend(widget.winfo_children())
        except Exception:
            pass
